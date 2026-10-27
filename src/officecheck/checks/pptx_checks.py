"""pptx checks."""

from __future__ import annotations

import math
from typing import Any

from pptx.enum.shapes import PP_PLACEHOLDER

from ..context import Context
from ..ooxml import parse_chart, q
from ..registry import (BOOL, COUNT_PROPS, COUNT_REQ, NONEMPTY_STR, POS_INT, TEXT_MATCH_PROPS, TEXT_OR_REGEX,
                        CheckError, Outcome, check)
from ..util import TextMatcher, count_expected, count_ok, first_n, norm_ws
from ._common import CHART_PROPS, match_charts

SLIDE = {"slide": POS_INT}
EMU_PER_PT = 12700


def _slides(ctx: Context, p: dict[str, Any]) -> list[tuple[int, Any]]:
    slides = list(ctx.pptx().slides)
    if "slide" in p:
        n = p["slide"]
        if n > len(slides):
            raise CheckError(f"presentation has {len(slides)} slide(s), no slide {n}")
        return [(n, slides[n - 1])]
    return list(enumerate(slides, 1))


def _walk_shapes(shapes: Any) -> Any:
    for sh in shapes:
        yield sh
        if sh.shape_type is not None and getattr(sh, "shapes", None) is not None and sh._element.tag == q("p:grpSp"):
            yield from _walk_shapes(sh.shapes)


@check("pptx.slide_count", COUNT_PROPS, **COUNT_REQ)
def slide_count(ctx: Context, p: dict[str, Any]) -> Outcome:
    n = len(ctx.pptx().slides)
    return Outcome(count_ok(n, p), count_expected(p), n)


def _title(slide: Any) -> str | None:
    t = slide.shapes.title
    if t is None or not t.has_text_frame:
        return None
    return norm_ws(t.text_frame.text)


@check("pptx.title_on_slide", {**SLIDE, **TEXT_MATCH_PROPS}, required=("slide",), **TEXT_OR_REGEX)
def title_on_slide(ctx: Context, p: dict[str, Any]) -> Outcome:
    (n, slide), = _slides(ctx, p)
    m = TextMatcher(p)
    t = _title(slide)
    return Outcome(m.matches(t), {"slide": n} | m.describe(), t)


@check("pptx.layout_on_slide", {**SLIDE, "layout": NONEMPTY_STR, "ignore_case": BOOL}, required=("slide", "layout"))
def layout_on_slide(ctx: Context, p: dict[str, Any]) -> Outcome:
    (n, slide), = _slides(ctx, p)
    name = slide.slide_layout.name
    want = p["layout"]
    ok = (name or "").casefold() == want.casefold() if p.get("ignore_case") else name == want
    return Outcome(ok, {"slide": n, "layout": want}, name)


def _notes(slide: Any) -> str:
    if not slide.has_notes_slide:
        return ""
    tf = slide.notes_slide.notes_text_frame
    return norm_ws(tf.text) if tf is not None else ""


@check("pptx.has_notes", {**SLIDE, "mode": {"enum": ["all", "any"]}, "min_chars": POS_INT})
def has_notes(ctx: Context, p: dict[str, Any]) -> Outcome:
    slides = _slides(ctx, p)
    need = p.get("min_chars", 1)
    have = {n: len(_notes(s)) for n, s in slides}
    with_notes = [n for n, c in have.items() if c >= need]
    mode = p.get("mode", "all")
    ok = bool(with_notes) if mode == "any" else (bool(have) and len(with_notes) == len(have))
    exp = ({"slide": p["slide"]} if "slide" in p else {"mode": mode}) | {"min_chars": need}
    return Outcome(ok, exp, {"slides_with_notes": with_notes,
                             "slides_without_notes": [n for n in have if n not in with_notes]})


@check("pptx.has_table", {**SLIDE, "min_rows": POS_INT, "min_cols": POS_INT, "min_count": POS_INT})
def has_table(ctx: Context, p: dict[str, Any]) -> Outcome:
    tables = []
    for n, slide in _slides(ctx, p):
        for sh in _walk_shapes(slide.shapes):
            if getattr(sh, "has_table", False) and sh.has_table:
                t = sh.table
                tables.append({"slide": n, "rows": len(t.rows), "cols": len(t.columns)})
    k = sum(1 for t in tables if t["rows"] >= p.get("min_rows", 1) and t["cols"] >= p.get("min_cols", 1))
    exp = {x: p[x] for x in ("slide", "min_rows", "min_cols", "min_count") if x in p} or {"min_count": 1}
    return Outcome(k >= p.get("min_count", 1), exp, {"matching": k, "tables": first_n(tables, 20)})


@check("pptx.has_chart", {**SLIDE, **CHART_PROPS})
def has_chart(ctx: Context, p: dict[str, Any]) -> Outcome:
    charts = []
    for n, slide in _slides(ctx, p):
        for sh in _walk_shapes(slide.shapes):
            if getattr(sh, "has_chart", False) and sh.has_chart:
                charts.append({"slide": n, **parse_chart(sh.chart._chartSpace)})
    return match_charts(charts, p)


def _image_count(slide: Any) -> int:
    root = slide._element.find(q("p:cSld"))
    tree = root.find(q("p:spTree")) if root is not None else None
    if tree is None:
        return 0
    n = 0
    for el in tree.iter():
        if el.tag == q("p:pic"):
            n += 1
        elif el.tag == q("p:sp"):
            sppr = el.find(q("p:spPr"))
            if sppr is not None and sppr.find(q("a:blipFill")) is not None:
                n += 1
    return n


@check("pptx.has_image", {**SLIDE, "min_count": POS_INT})
def has_image(ctx: Context, p: dict[str, Any]) -> Outcome:
    per = {n: _image_count(s) for n, s in _slides(ctx, p)}
    total = sum(per.values())
    exp = {x: p[x] for x in ("slide", "min_count") if x in p} or {"min_count": 1}
    return Outcome(total >= p.get("min_count", 1), exp, {"images": total, "per_slide": per})


# ---- text_fits -------------------------------------------------------------------------------

_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
_BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE,
               PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.VERTICAL_OBJECT}


def _lst_size(el: Any, level: int) -> float | None:
    """defRPr sz (points) for paragraph ``level`` from an a:lstStyle / p:*Style element."""
    if el is None:
        return None
    lvl = el.find(q(f"a:lvl{level + 1}pPr"))
    if lvl is not None:
        d = lvl.find(q("a:defRPr"))
        if d is not None and d.get("sz"):
            return int(d.get("sz")) / 100.0
    return None


def _shape_lst(sp_el: Any) -> Any:
    body = sp_el.find(q("p:txBody"))
    return body.find(q("a:lstStyle")) if body is not None else None


def _inherited_size(shape: Any, slide: Any, level: int, default_pt: float) -> float:
    sz = _lst_size(_shape_lst(shape._element), level)
    if sz:
        return sz
    if not shape.is_placeholder:
        return default_pt
    pf = shape.placeholder_format
    layout = slide.slide_layout
    master = layout.slide_master
    for container in (layout, master):
        for ph in container.placeholders:
            ppf = ph.placeholder_format
            same = (ppf.idx == pf.idx) if container is layout else (ppf.type == pf.type or
                                                                      (pf.type in _BODY_TYPES and ppf.type == PP_PLACEHOLDER.BODY))
            if same:
                sz = _lst_size(_shape_lst(ph._element), level)
                if sz:
                    return sz
                break
    styles = master._element.find(q("p:txStyles"))
    if styles is not None:
        name = "p:titleStyle" if pf.type in _TITLE_TYPES else ("p:bodyStyle" if pf.type in _BODY_TYPES else "p:otherStyle")
        sz = _lst_size(styles.find(q(name)), level)
        if sz:
            return sz
    return default_pt


def _fit(shape: Any, slide: Any, p: dict[str, Any]) -> dict[str, Any] | None:
    if not shape.has_text_frame:
        return None
    tf = shape.text_frame
    if not norm_ws(tf.text):
        return None
    w, h = shape.width, shape.height
    if not w or not h:
        return None
    body_pr = tf._txBody.find(q("a:bodyPr"))
    ins = {k: int(body_pr.get(k)) if body_pr is not None and body_pr.get(k) else d
           for k, d in (("lIns", 91440), ("rIns", 91440), ("tIns", 45720), ("bIns", 45720))}
    vertical = body_pr is not None and (body_pr.get("vert") or "horz") not in ("horz",)
    usable_w = (w - ins["lIns"] - ins["rIns"]) / EMU_PER_PT
    usable_h = (h - ins["tIns"] - ins["bIns"]) / EMU_PER_PT
    if vertical:
        usable_w, usable_h = usable_h, usable_w
    if usable_w <= 0 or usable_h <= 0:
        return {"shape": shape.name, "fill": math.inf}
    wrap = body_pr is None or body_pr.get("wrap") != "none"
    scale, spacing = 1.0, 1.0
    norm = body_pr.find(q("a:normAutofit")) if body_pr is not None else None
    if norm is not None:
        scale = int(norm.get("fontScale", "100000")) / 100000.0
        spacing = 1.0 - int(norm.get("lnSpcReduction", "0")) / 100000.0
    cw, lh = p.get("char_width_em", 0.5), p.get("line_height_em", 1.2)
    default_pt = p.get("default_font_pt", 18.0)
    need_h, chars, widest = 0.0, 0, 0.0
    for para in tf.paragraphs:
        base = _inherited_size(shape, slide, para.level, default_pt)
        line_w, max_fs, lines = 0.0, 0.0, 0
        segs: list[tuple[str, float]] = [(r.text, r.font.size.pt if r.font.size else base) for r in para.runs]
        if not segs:
            segs = [("", base)]
        for text, fs in segs:
            fs *= scale
            max_fs = max(max_fs, fs)
            for i, piece in enumerate(text.split("\v")):
                if i:  # explicit line break
                    lines += max(1, math.ceil(line_w / usable_w)) if wrap else 1
                    widest = max(widest, line_w)
                    line_w = 0.0
                line_w += len(piece) * cw * fs
                chars += len(piece)
        lines += max(1, math.ceil(line_w / usable_w)) if wrap else 1
        widest = max(widest, line_w)
        need_h += lines * lh * max_fs * spacing
    fill = need_h / usable_h
    if not wrap:
        fill = max(fill, widest / usable_w)
    return {"shape": shape.name, "chars": chars, "fill": round(fill, 3)}


@check("pptx.text_fits", {**SLIDE, "max_fill": {"type": "number", "exclusiveMinimum": 0},
                          "char_width_em": {"type": "number", "exclusiveMinimum": 0},
                          "line_height_em": {"type": "number", "exclusiveMinimum": 0},
                          "default_font_pt": {"type": "number", "exclusiveMinimum": 0},
                          "placeholders_only": BOOL})
def text_fits(ctx: Context, p: dict[str, Any]) -> Outcome:
    max_fill = p.get("max_fill", 1.0)
    over, checked = [], 0
    for n, slide in _slides(ctx, p):
        for sh in _walk_shapes(slide.shapes):
            if p.get("placeholders_only") and not sh.is_placeholder:
                continue
            r = _fit(sh, slide, p)
            if r is None:
                continue
            checked += 1
            if r["fill"] > max_fill:
                over.append({"slide": n, **r})
    return Outcome(not over, {"max_fill": max_fill},
                   {"text_shapes": checked, "overflowing": first_n(over, 20)})
