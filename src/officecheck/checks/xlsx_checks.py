"""xlsx checks."""

from __future__ import annotations

import datetime as _dt
import math
import re
from typing import Any

from openpyxl.utils.cell import (column_index_from_string, coordinate_from_string, get_column_letter,
                                 range_boundaries)

from ..context import Context
from ..formula import (ERROR_CODES, ExcelError, UnsupportedFormula, WorkbookEvaluator, cell_formula,
                       compare)
from ..ooxml import iter_charts_in_xlsx, parse_chart, read_xml
from ._common import CHART_PROPS, match_charts
from ..registry import (BOOL, COUNT_PROPS, COUNT_REQ, NONEMPTY_STR, NONNEG_INT, POS_INT,
                        CheckError, Outcome, check)
from ..util import count_expected, count_ok, first_n, norm_ws

MAX_ROW, MAX_COL = 1_048_576, 16_384
SHEET = {"sheet": NONEMPTY_STR}
CELL = {"type": "string", "pattern": r"^(?:(?:'(?:[^']|'')+'|[^!]+)!)?\$?[A-Za-z]{1,3}\$?[0-9]+$"}
RANGE = {"type": "string", "minLength": 2}
COLUMN_LETTER = {"type": "string", "pattern": "^[A-Za-z]{1,3}$"}


# ---- helpers ---------------------------------------------------------------------------------


def _sheet_name(ctx: Context, name: str | None) -> str:
    wb = ctx.xlsx(False)
    if name is None:
        return wb.sheetnames[0]
    if name in wb.sheetnames:
        return name
    for s in wb.sheetnames:
        if s.casefold() == name.casefold():
            return s
    raise CheckError(f"no sheet named {name!r} (sheets: {wb.sheetnames})")


def _split_cell(ctx: Context, p: dict[str, Any]) -> tuple[str, int, int, str]:
    ref = p["cell"]
    sheet = p.get("sheet")
    if "!" in ref:
        sheet, ref = ref.rsplit("!", 1)
        if sheet.startswith("'"):
            sheet = sheet[1:-1].replace("''", "'")
    sheet = _sheet_name(ctx, sheet)
    col, row = coordinate_from_string(ref.replace("$", "").upper())
    return sheet, row, column_index_from_string(col), f"{col}{row}"


def _evaluator(ctx: Context, prefer_cached: bool = True) -> WorkbookEvaluator:
    key = f"xlsx_eval_{prefer_cached}"
    if key not in ctx._cache:
        ctx._cache[key] = WorkbookEvaluator(ctx.xlsx(False), ctx.xlsx(True), prefer_cached)
    return ctx._cache[key]


def _value(ctx: Context, sheet: str, row: int, col: int) -> Any:
    """Cached value of a cell, or its evaluated formula when no cached value is stored."""
    try:
        return _evaluator(ctx)._get(sheet, row, col)
    except UnsupportedFormula as e:
        raise CheckError(f"{sheet}!{get_column_letter(col)}{row}: {e}") from e


def _display(v: Any) -> Any:
    if isinstance(v, ExcelError):
        return v.code
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return int(v)
    return v


def _header_col(ctx: Context, sheet: str, p: dict[str, Any], key: str = "column") -> int:
    if f"{key}_letter" in p:
        return column_index_from_string(p[f"{key}_letter"].upper())
    want = p[key]
    ws = ctx.xlsx(False)[sheet]
    row = p.get("header_row", 1)
    fold = (lambda s: s.casefold()) if p.get("ignore_case") else (lambda s: s)
    headers = []
    for c in range(1, (ws.max_column or 0) + 1):
        v = _value(ctx, sheet, row, c)
        if v is None:
            continue
        h = norm_ws(str(_display(v)))
        headers.append(h)
        if fold(h) == fold(norm_ws(want)):
            return c
    raise CheckError(f"no column with header {want!r} in row {row} of {sheet!r} (headers: {headers})")


def _rects(rng: str) -> list[tuple[int, int, int, int]]:
    """'A1:B3 D5' -> [(min_col, min_row, max_col, max_row), ...] with open ends filled in."""
    out = []
    for part in re.split(r"[\s,]+", rng.replace("$", "").strip()):
        if not part:
            continue
        if "!" in part:
            part = part.rsplit("!", 1)[1]
        c1, r1, c2, r2 = range_boundaries(part.upper())
        out.append((c1 or 1, r1 or 1, c2 or MAX_COL, r2 or MAX_ROW))
    return out


def _covered(target: tuple[int, int, int, int], rects: list[tuple[int, int, int, int]]) -> bool:
    """True if every cell of ``target`` lies inside the union of ``rects``."""
    tc1, tr1, tc2, tr2 = target
    cuts = {tr1, tr2 + 1}
    for _c1, r1, _c2, r2 in rects:
        for r in (r1, r2 + 1):
            if tr1 < r <= tr2:
                cuts.add(r)
    bounds = sorted(cuts)
    for lo in bounds[:-1]:  # each row segment [lo, next) has a constant set of covering rects
        spans = sorted((max(c1, tc1), min(c2, tc2)) for c1, r1, c2, r2 in rects
                       if r1 <= lo <= r2 and c2 >= tc1 and c1 <= tc2)
        nxt = tc1
        for a, b in spans:
            if a > nxt:
                break
            nxt = max(nxt, b + 1)
        if nxt <= tc2:
            return False
    return True


def _intersects(target: tuple[int, int, int, int], rects: list[tuple[int, int, int, int]]) -> bool:
    tc1, tr1, tc2, tr2 = target
    return any(c1 <= tc2 and c2 >= tc1 and r1 <= tr2 and r2 >= tr1 for c1, r1, c2, r2 in rects)


def _range_ok(target: str, ranges: list[str], mode: str) -> bool:
    rects = [r for s in ranges for r in _rects(s)]
    targets = _rects(target)
    if mode == "intersects":
        return any(_intersects(t, rects) for t in targets)
    return all(_covered(t, rects) for t in targets)


# ---- checks ----------------------------------------------------------------------------------


@check("xlsx.sheet_exists", {"sheet": NONEMPTY_STR, "ignore_case": BOOL}, required=("sheet",))
def sheet_exists(ctx: Context, p: dict[str, Any]) -> Outcome:
    names = ctx.xlsx(False).sheetnames
    want = p["sheet"]
    ok = want in names or (p.get("ignore_case", False) and want.casefold() in {n.casefold() for n in names})
    return Outcome(ok, want, names)


def _equal(expected: Any, actual: Any, p: dict[str, Any]) -> bool:
    if isinstance(actual, ExcelError):
        actual = actual.code
    if expected is None:
        return actual is None or actual == ""
    if isinstance(expected, bool):
        return isinstance(actual, bool) and actual == expected
    if isinstance(expected, (int, float)):
        if isinstance(actual, bool) or not isinstance(actual, (int, float)):
            return False
        tol = p.get("tolerance", 1e-9)
        return math.isclose(float(actual), float(expected), rel_tol=1e-12, abs_tol=tol)
    if isinstance(expected, str):
        if isinstance(actual, (_dt.datetime, _dt.date, _dt.time)):
            iso = actual.isoformat()
            return iso == expected or (isinstance(actual, _dt.datetime) and actual.time() == _dt.time(0)
                                       and actual.date().isoformat() == expected)
        if not isinstance(actual, str):
            return False
        if p.get("ignore_case"):
            return actual.casefold() == expected.casefold()
        return actual == expected
    return False


@check("xlsx.cell_equals", {**SHEET, "cell": CELL, "value": {"type": ["string", "number", "boolean", "null"]},
                            "tolerance": {"type": "number", "minimum": 0}, "ignore_case": BOOL,
                            "source": {"enum": ["auto", "cached", "evaluate"]}},
       required=("cell", "value"))
def cell_equals(ctx: Context, p: dict[str, Any]) -> Outcome:
    sheet, row, col, coord = _split_cell(ctx, p)
    fcell = ctx.xlsx(False)[sheet].cell(row=row, column=col)
    formula = cell_formula(fcell)
    source = p.get("source", "auto")
    exp = {"cell": f"{sheet}!{coord}", "value": p["value"]}
    if formula is None:
        actual, how = fcell.value, "literal"
    else:
        cached = ctx.xlsx(True)[sheet].cell(row=row, column=col).value
        if source == "cached" or (source == "auto" and cached is not None):
            actual, how = cached, "cached"
        else:
            try:
                ev = _evaluator(ctx, prefer_cached=(source == "auto"))
                actual, how = _display(ev.evaluate(sheet, row, col)), "evaluated"
            except UnsupportedFormula as e:
                return Outcome(False, exp, {"formula": formula, "error": f"cannot evaluate: {e}"})
    act: dict[str, Any] = {"value": _display(actual), "source": how}
    if formula is not None:
        act["formula"] = formula
    return Outcome(_equal(p["value"], actual, p), exp, act)


@check("xlsx.cell_has_formula", {**SHEET, "cell": CELL, "pattern": NONEMPTY_STR, "ignore_case": BOOL},
       required=("cell",))
def cell_has_formula(ctx: Context, p: dict[str, Any]) -> Outcome:
    sheet, row, col, coord = _split_cell(ctx, p)
    f = cell_formula(ctx.xlsx(False)[sheet].cell(row=row, column=col))
    exp: dict[str, Any] = {"cell": f"{sheet}!{coord}", "formula": True}
    ok = f is not None
    if "pattern" in p:
        exp["pattern"] = p["pattern"]
        flags = re.IGNORECASE if p.get("ignore_case", True) else 0
        ok = ok and re.search(p["pattern"], f, flags) is not None
    return Outcome(ok, exp, f)


@check("xlsx.column_exists", {**SHEET, "header": NONEMPTY_STR, "header_row": POS_INT, "ignore_case": BOOL},
       required=("header",))
def column_exists(ctx: Context, p: dict[str, Any]) -> Outcome:
    sheet = _sheet_name(ctx, p.get("sheet"))
    try:
        col = _header_col(ctx, sheet, {**p, "column": p["header"]})
        return Outcome(True, p["header"], {"column": get_column_letter(col)})
    except CheckError as e:
        return Outcome(False, p["header"], str(e))


def _nonempty(v: Any) -> bool:
    return v is not None and not (isinstance(v, str) and v.strip() == "")


@check("xlsx.row_count", {**SHEET, **COUNT_PROPS, "header_rows": NONNEG_INT}, **COUNT_REQ)
def row_count(ctx: Context, p: dict[str, Any]) -> Outcome:
    sheet = _sheet_name(ctx, p.get("sheet"))
    ws = ctx.xlsx(False)[sheet]
    hdr = p.get("header_rows", 1)
    n = sum(1 for r in ws.iter_rows(min_row=hdr + 1, values_only=True) if any(_nonempty(v) for v in r))
    return Outcome(count_ok(n, p), count_expected(p) | {"header_rows": hdr}, n)


@check("xlsx.chart_exists", {**SHEET, **CHART_PROPS})
def chart_exists(ctx: Context, p: dict[str, Any]) -> Outcome:
    zf = ctx.zip()
    want_sheet = _sheet_name(ctx, p["sheet"]) if "sheet" in p else None
    charts = []
    for sheet, part in iter_charts_in_xlsx(zf):
        if want_sheet is None or sheet == want_sheet:
            charts.append({"sheet": sheet, **parse_chart(read_xml(zf, part))})
    return match_charts(charts, p)


@check("xlsx.conditional_format_on", {**SHEET, "range": RANGE, "rule_type": NONEMPTY_STR,
                                       "mode": {"enum": ["covers", "intersects"]}}, required=("range",))
def conditional_format_on(ctx: Context, p: dict[str, Any]) -> Outcome:
    sheet = _sheet_name(ctx, p.get("sheet"))
    ws = ctx.xlsx(False)[sheet]
    found = []
    for cf in ws.conditional_formatting:
        types = sorted({r.type for r in cf.rules if r.type})
        found.append({"range": str(cf.sqref), "rule_types": types})
    ranges = [f["range"] for f in found if "rule_type" not in p or p["rule_type"] in f["rule_types"]]
    ok = bool(ranges) and _range_ok(p["range"], ranges, p.get("mode", "covers"))
    exp = {"range": p["range"], "mode": p.get("mode", "covers")} | ({"rule_type": p["rule_type"]} if "rule_type" in p else {})
    return Outcome(ok, exp, found)


@check("xlsx.data_validation_on", {**SHEET, "range": RANGE, "validation_type": {"enum": [
    "whole", "decimal", "list", "date", "time", "textLength", "custom"]}, "mode": {"enum": ["covers", "intersects"]}},
       required=("range",))
def data_validation_on(ctx: Context, p: dict[str, Any]) -> Outcome:
    sheet = _sheet_name(ctx, p.get("sheet"))
    ws = ctx.xlsx(False)[sheet]
    found = []
    for dv in ws.data_validations.dataValidation:
        found.append({"range": str(dv.sqref), "type": dv.type, "formula1": dv.formula1})
    ranges = [f["range"] for f in found if "validation_type" not in p or f["type"] == p["validation_type"]]
    ok = bool(ranges) and _range_ok(p["range"], ranges, p.get("mode", "covers"))
    exp = {"range": p["range"], "mode": p.get("mode", "covers")}
    if "validation_type" in p:
        exp["validation_type"] = p["validation_type"]
    return Outcome(ok, exp, found)


@check("xlsx.no_error_values", {**SHEET, "errors": {"type": "array", "items": {"enum": list(ERROR_CODES)}, "minItems": 1},
                                "evaluate_uncached": BOOL})
def no_error_values(ctx: Context, p: dict[str, Any]) -> Outcome:
    errors = set(p.get("errors", ERROR_CODES))
    vwb, fwb = ctx.xlsx(True), ctx.xlsx(False)
    sheets = [_sheet_name(ctx, p["sheet"])] if "sheet" in p else vwb.sheetnames
    hits, uncached, unevaluable = [], 0, 0
    for s in sheets:
        fws = fwb[s]
        for row in vwb[s].iter_rows():
            for c in row:
                if c.data_type == "e" and c.value in errors:
                    hits.append(f"{s}!{c.coordinate}={c.value}")
                elif c.value is None and cell_formula(fws[c.coordinate]) is not None:
                    uncached += 1
                    if p.get("evaluate_uncached"):
                        try:
                            v = _evaluator(ctx).evaluate(s, c.row, c.column)
                        except UnsupportedFormula:
                            unevaluable += 1
                            continue
                        if isinstance(v, ExcelError) and v.code in errors:
                            hits.append(f"{s}!{c.coordinate}={v.code} (evaluated)")
    act: dict[str, Any] = {"errors": first_n(hits, 50), "count": len(hits), "formulas_without_cached_value": uncached}
    if p.get("evaluate_uncached"):
        act["unevaluable"] = unevaluable
    return Outcome(not hits, {"errors": sorted(errors), "count": 0}, act)


def _column_values(ctx: Context, sheet: str, col: int, header_row: int) -> list[tuple[int, Any]]:
    ws = ctx.xlsx(False)[sheet]
    return [(r, _value(ctx, sheet, r, col)) for r in range(header_row + 1, (ws.max_row or 0) + 1)]


@check("xlsx.sorted_by", {**SHEET, "column": NONEMPTY_STR, "column_letter": COLUMN_LETTER,
                          "order": {"enum": ["asc", "desc"]}, "header_row": NONNEG_INT, "ignore_case": BOOL},
       oneOf=[{"required": ["column"]}, {"required": ["column_letter"]}])
def sorted_by(ctx: Context, p: dict[str, Any]) -> Outcome:
    sheet = _sheet_name(ctx, p.get("sheet"))
    col = _header_col(ctx, sheet, p)
    order = p.get("order", "asc")
    vals = [(r, v) for r, v in _column_values(ctx, sheet, col, p.get("header_row", 1)) if v is not None]
    bad = None
    for (_r0, a), (r1, b) in zip(vals, vals[1:], strict=False):
        c = compare(a, b) if not isinstance(a, ExcelError) and not isinstance(b, ExcelError) else None
        if c is None or (c > 0 if order == "asc" else c < 0):
            bad = {"row": r1, "previous": _display(a), "value": _display(b)}
            break
    exp = {"column": p.get("column") or p.get("column_letter"), "order": order}
    act = {"column": get_column_letter(col), "values": len(vals)}
    if bad:
        act["first_out_of_order"] = bad
    return Outcome(bad is None and bool(vals), exp, act)


@check("xlsx.no_duplicates", {**SHEET, "columns": {"type": "array", "items": NONEMPTY_STR, "minItems": 1},
                              "column_letters": {"type": "array", "items": COLUMN_LETTER, "minItems": 1},
                              "header_row": NONNEG_INT, "ignore_case": BOOL},
       oneOf=[{"required": ["columns"]}, {"required": ["column_letters"]}])
def no_duplicates(ctx: Context, p: dict[str, Any]) -> Outcome:
    sheet = _sheet_name(ctx, p.get("sheet"))
    if "column_letters" in p:
        cols = [column_index_from_string(c.upper()) for c in p["column_letters"]]
    else:
        cols = [_header_col(ctx, sheet, {**p, "column": c}) for c in p["columns"]]
    hr = p.get("header_row", 1)
    ws = ctx.xlsx(False)[sheet]
    seen: dict[tuple[Any, ...], int] = {}
    dups = []
    for r in range(hr + 1, (ws.max_row or 0) + 1):
        key = tuple(_value(ctx, sheet, r, c) for c in cols)
        if all(k is None for k in key):
            continue
        norm = tuple((k.casefold() if p.get("ignore_case") else k) if isinstance(k, str)
                     else (k.code if isinstance(k, ExcelError) else k) for k in key)
        if norm in seen:
            dups.append({"row": r, "duplicate_of_row": seen[norm], "key": [_display(k) for k in key]})
        else:
            seen[norm] = r
    exp = {"columns": p.get("columns") or p.get("column_letters"), "duplicates": 0}
    return Outcome(not dups, exp, {"duplicates": len(dups), "examples": first_n(dups, 10)})


@check("xlsx.freeze_panes", {**SHEET, "cell": {"oneOf": [CELL, {"type": "null"}]}, "rows": NONNEG_INT,
                             "cols": NONNEG_INT},
       oneOf=[{"required": ["cell"], "not": {"anyOf": [{"required": ["rows"]}, {"required": ["cols"]}]}},
              {"anyOf": [{"required": ["rows"]}, {"required": ["cols"]}], "not": {"required": ["cell"]}}])
def freeze_panes(ctx: Context, p: dict[str, Any]) -> Outcome:
    sheet = _sheet_name(ctx, p.get("sheet"))
    ws = ctx.xlsx(False)[sheet]
    pane = ws.sheet_view.pane
    actual = pane.topLeftCell if pane is not None and pane.state in ("frozen", "frozenSplit") else None
    if actual is None and pane is not None and pane.state in ("frozen", "frozenSplit"):
        x, y = int(pane.xSplit or 0), int(pane.ySplit or 0)
        actual = f"{get_column_letter(x + 1)}{y + 1}" if (x or y) else None
    if "cell" in p:
        want = p["cell"].replace("$", "").upper() if p["cell"] else None
    else:
        want = f"{get_column_letter(p.get('cols', 0) + 1)}{p.get('rows', 0) + 1}"
        if want == "A1":
            want = None
    return Outcome(actual == want, {"sheet": sheet, "frozen_at": want}, {"frozen_at": actual})
