"""A small, deterministic evaluator for a subset of Excel formulas.

Supported: numbers, strings, booleans, error literals, cell refs, ranges (A1:B3, A:A, with
optional sheet qualifiers and $), arithmetic (+ - * / ^ %), concatenation (&), comparisons
(= <> < > <= >=) and the functions SUM, AVERAGE, COUNT, COUNTA, MIN, MAX, IF, IFERROR,
VLOOKUP, ROUND, ABS, AND, OR, NOT, CONCATENATE/CONCAT.

Anything else raises ``UnsupportedFormula`` so a grader never reports a guessed value.
Excel errors are values (``ExcelError``), propagated like Excel does.
"""

from __future__ import annotations

import datetime as _dt
import decimal
import math
import re
from dataclasses import dataclass
from typing import Any, Callable

from openpyxl.utils.cell import column_index_from_string, coordinate_from_string


class UnsupportedFormula(Exception):
    pass


@dataclass(frozen=True)
class ExcelError:
    code: str

    def __str__(self) -> str:
        return self.code


DIV0, NA, REF, VALUE, NAME, NUM = (ExcelError(c) for c in ("#DIV/0!", "#N/A", "#REF!", "#VALUE!", "#NAME?", "#NUM!"))
ERROR_CODES = ("#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A")


@dataclass
class Range:
    rows: list[list[Any]]

    def flat(self) -> list[Any]:
        return [v for r in self.rows for v in r]


# ---- tokenizer -------------------------------------------------------------------------------

_SHEET = r"(?:'(?:[^']|'')+'|[A-Za-z_][\w.]*)!"
_CELL = r"\$?[A-Za-z]{1,3}\$?\d+"
_TOKEN = re.compile(
    rf"""
    (?P<ws>\s+)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<err>\#(?:DIV/0!|N/A|REF!|VALUE!|NAME\?|NUM!|NULL!))
  | (?P<ref>(?:{_SHEET})?(?:{_CELL}(?::{_CELL})?|\$?[A-Za-z]{{1,3}}:\$?[A-Za-z]{{1,3}}))(?![\w(])
  | (?P<num>(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?)
  | (?P<func>[A-Za-z_][\w.]*)\s*\(
  | (?P<bool>TRUE|FALSE)(?![\w(])
  | (?P<name>[A-Za-z_][\w.]*)
  | (?P<op><>|<=|>=|[-+*/^&%=<>(),;])
    """,
    re.VERBOSE | re.IGNORECASE,
)


def tokenize(src: str) -> list[tuple[str, str]]:
    out, pos = [], 0
    while pos < len(src):
        m = _TOKEN.match(src, pos)
        if not m:
            raise UnsupportedFormula(f"cannot parse formula at {src[pos:pos + 10]!r}")
        pos = m.end()
        kind = m.lastgroup
        if kind == "ws":
            continue
        if kind == "name":
            raise UnsupportedFormula(f"defined names are not supported: {m.group()}")
        out.append((kind, m.group(kind)))
    return out


# ---- parser (to tuples) ----------------------------------------------------------------------


class _Parser:
    def __init__(self, tokens: list[tuple[str, str]]):
        self.t, self.i = tokens, 0

    def peek(self) -> tuple[str, str] | None:
        return self.t[self.i] if self.i < len(self.t) else None

    def take(self) -> tuple[str, str]:
        tok = self.peek()
        if tok is None:
            raise UnsupportedFormula("unexpected end of formula")
        self.i += 1
        return tok

    def accept(self, *ops: str) -> str | None:
        tok = self.peek()
        if tok and tok[0] == "op" and tok[1] in ops:
            self.i += 1
            return tok[1]
        return None

    def parse(self) -> Any:
        node = self.comparison()
        if self.peek() is not None:
            raise UnsupportedFormula(f"unexpected token {self.peek()[1]!r}")
        return node

    def comparison(self) -> Any:
        node = self.concat()
        while (op := self.accept("=", "<>", "<", ">", "<=", ">=")):
            node = ("bin", op, node, self.concat())
        return node

    def concat(self) -> Any:
        node = self.additive()
        while self.accept("&"):
            node = ("bin", "&", node, self.additive())
        return node

    def additive(self) -> Any:
        node = self.mult()
        while (op := self.accept("+", "-")):
            node = ("bin", op, node, self.mult())
        return node

    def mult(self) -> Any:
        node = self.power()
        while (op := self.accept("*", "/")):
            node = ("bin", op, node, self.power())
        return node

    def power(self) -> Any:
        node = self.unary()
        while self.accept("^"):
            node = ("bin", "^", node, self.unary())
        return node

    def unary(self) -> Any:  # in Excel, negation binds tighter than ^ (=-2^2 is 4)
        if (op := self.accept("-", "+")):
            inner = self.unary()
            return ("neg", inner) if op == "-" else ("pos", inner)
        node = self.primary()
        while self.accept("%"):
            node = ("pct", node)
        return node

    def primary(self) -> Any:
        kind, val = self.take()
        if kind == "num":
            return ("lit", float(val))
        if kind == "str":
            return ("lit", val[1:-1].replace('""', '"'))
        if kind == "bool":
            return ("lit", val.upper() == "TRUE")
        if kind == "err":
            return ("lit", ExcelError(val.upper()))
        if kind == "ref":
            return ("ref", val)
        if kind == "func":
            args: list[Any] = []
            if not self.accept(")"):
                while True:
                    tok = self.peek()
                    if tok and tok[0] == "op" and tok[1] in (",", ";", ")"):
                        args.append(("lit", None))  # omitted argument
                    else:
                        args.append(self.comparison())
                    if self.accept(")"):
                        break
                    if not self.accept(",", ";"):
                        raise UnsupportedFormula("expected , or ) in function call")
            return ("call", val.upper(), args)
        if kind == "op" and val == "(":
            node = self.comparison()
            if not self.accept(")"):
                raise UnsupportedFormula("missing )")
            return node
        raise UnsupportedFormula(f"unexpected token {val!r}")


def parse(formula: str) -> Any:
    src = formula[1:] if formula.startswith("=") else formula
    return _Parser(tokenize(src)).parse()


# ---- coercion helpers ------------------------------------------------------------------------


def _is_num(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def to_number(v: Any) -> float | ExcelError:
    if isinstance(v, ExcelError):
        return v
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if v is None:
        return 0.0
    if _is_num(v):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip().replace(",", "")) if v.strip() else VALUE
        except ValueError:
            return VALUE
    return VALUE


def to_bool(v: Any) -> bool | ExcelError:
    if isinstance(v, ExcelError):
        return v
    if isinstance(v, bool):
        return v
    if v is None:
        return False
    if _is_num(v):
        return v != 0
    if isinstance(v, str):
        if v.upper() in ("TRUE", "FALSE"):
            return v.upper() == "TRUE"
        return VALUE
    return VALUE


def fmt_number(x: float) -> str:
    if x == int(x) and abs(x) < 1e15:
        return str(int(x))
    return format(x, ".15g")


def to_text(v: Any) -> str | ExcelError:
    if isinstance(v, ExcelError):
        return v
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if _is_num(v):
        return fmt_number(float(v))
    return str(v)


def _type_rank(v: Any) -> int:
    if isinstance(v, bool):
        return 2
    if isinstance(v, str):
        return 1
    return 0


def compare(a: Any, b: Any) -> int:
    """Excel ordering: numbers < text < booleans; text compares case-insensitively."""
    if a is None:
        a = "" if isinstance(b, str) else (False if isinstance(b, bool) else 0.0)
    if b is None:
        b = "" if isinstance(a, str) else (False if isinstance(a, bool) else 0.0)
    ra, rb = _type_rank(a), _type_rank(b)
    if ra != rb:
        return -1 if ra < rb else 1
    if ra == 1:
        a, b = a.casefold(), b.casefold()
    return (a > b) - (a < b)


def round_half_away(x: float, n: int) -> float:
    q = decimal.Decimal(1).scaleb(-n)
    return float(decimal.Decimal(repr(x)).quantize(q, rounding=decimal.ROUND_HALF_UP))


# ---- evaluator -------------------------------------------------------------------------------

Getter = Callable[[str, int, int], Any]  # (sheet, row, col) -> value (already evaluated)


class Evaluator:
    """Evaluate formulas against a workbook through a value getter."""

    def __init__(self, get: Getter, sheet_names: list[str], max_rows: Callable[[str], int]):
        self.get = get
        self.sheets = {s.casefold(): s for s in sheet_names}
        self.max_rows = max_rows

    def eval_formula(self, formula: str, sheet: str) -> Any:
        v = self.ev(parse(formula), sheet)
        if isinstance(v, Range):
            v = self._scalar(v)
        return v

    # -- references --
    def _resolve_ref(self, ref: str, sheet: str) -> Range:
        if "!" in ref:
            sname, ref = ref.rsplit("!", 1)
            if sname.startswith("'"):
                sname = sname[1:-1].replace("''", "'")
            real = self.sheets.get(sname.casefold())
            if real is None:
                return Range([[REF]])
            sheet = real
        ref = ref.replace("$", "").upper()
        if ":" in ref:
            a, b = ref.split(":")
            if a.isalpha():  # whole columns
                c1, c2 = sorted((column_index_from_string(a), column_index_from_string(b)))
                r1, r2 = 1, max(1, self.max_rows(sheet))
            else:
                (ca, ra), (cb, rb) = coordinate_from_string(a), coordinate_from_string(b)
                c1, c2 = sorted((column_index_from_string(ca), column_index_from_string(cb)))
                r1, r2 = sorted((ra, rb))
        else:
            col, row = coordinate_from_string(ref)
            c1 = c2 = column_index_from_string(col)
            r1 = r2 = row
        if (r2 - r1 + 1) * (c2 - c1 + 1) > 2_000_000:
            raise UnsupportedFormula("range too large")
        return Range([[self.get(sheet, r, c) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)])

    @staticmethod
    def _scalar(v: Any) -> Any:
        if isinstance(v, Range):
            flat = v.flat()
            return flat[0] if len(flat) == 1 else VALUE
        return v

    # -- core --
    def ev(self, node: Any, sheet: str) -> Any:
        kind = node[0]
        if kind == "lit":
            return node[1]
        if kind == "ref":
            return self._resolve_ref(node[1], sheet)
        if kind == "neg":
            x = to_number(self._scalar(self.ev(node[1], sheet)))
            return x if isinstance(x, ExcelError) else -x
        if kind == "pos":
            return self._scalar(self.ev(node[1], sheet))
        if kind == "pct":
            x = to_number(self._scalar(self.ev(node[1], sheet)))
            return x if isinstance(x, ExcelError) else x / 100.0
        if kind == "bin":
            return self._binary(node[1], self._scalar(self.ev(node[2], sheet)),
                                self._scalar(self.ev(node[3], sheet)))
        if kind == "call":
            return self._call(node[1], node[2], sheet)
        raise UnsupportedFormula(f"bad node {kind}")

    def _binary(self, op: str, a: Any, b: Any) -> Any:
        if op in ("=", "<>", "<", ">", "<=", ">="):
            for v in (a, b):
                if isinstance(v, ExcelError):
                    return v
            c = compare(a, b)
            return {"=": c == 0, "<>": c != 0, "<": c < 0, ">": c > 0, "<=": c <= 0, ">=": c >= 0}[op]
        if op == "&":
            ta, tb = to_text(a), to_text(b)
            for v in (ta, tb):
                if isinstance(v, ExcelError):
                    return v
            return ta + tb
        x, y = to_number(a), to_number(b)
        for v in (x, y):
            if isinstance(v, ExcelError):
                return v
        if op == "+":
            return x + y
        if op == "-":
            return x - y
        if op == "*":
            return x * y
        if op == "/":
            return DIV0 if y == 0 else x / y
        if op == "^":
            if x == 0 and y < 0:
                return DIV0
            try:
                r = x ** y
            except (OverflowError, ZeroDivisionError):
                return NUM
            return NUM if isinstance(r, complex) or not math.isfinite(r) else r
        raise UnsupportedFormula(f"operator {op}")

    # -- functions --
    def _numbers(self, args: list[Any], sheet: str, *, count_mode: bool = False) -> list[float] | ExcelError:
        """Numbers per SUM/AVERAGE/MIN/MAX rules: in ranges only real numbers count."""
        out: list[float] = []
        for a in args:
            v = self.ev(a, sheet)
            if isinstance(v, Range):
                for x in v.flat():
                    if isinstance(x, ExcelError):
                        if count_mode:
                            continue
                        return x
                    if _is_num(x):
                        out.append(float(x))
            else:
                if v is None:
                    continue
                if count_mode:  # COUNT: literal numbers, booleans and numeric text count
                    n = to_number(v)
                    if not isinstance(n, ExcelError):
                        out.append(n)
                    continue
                n = to_number(v)
                if isinstance(n, ExcelError):
                    return n
                out.append(n)
        return out

    def _call(self, name: str, args: list[Any], sheet: str) -> Any:
        if name in ("SUM", "AVERAGE", "MIN", "MAX"):
            nums = self._numbers(args, sheet)
            if isinstance(nums, ExcelError):
                return nums
            if name == "SUM":
                return math.fsum(nums)
            if name == "AVERAGE":
                return DIV0 if not nums else math.fsum(nums) / len(nums)
            if not nums:
                return 0.0
            return min(nums) if name == "MIN" else max(nums)
        if name == "COUNT":
            nums = self._numbers(args, sheet, count_mode=True)
            return float(len(nums))
        if name == "COUNTA":
            n = 0
            for a in args:
                v = self.ev(a, sheet)
                vals = v.flat() if isinstance(v, Range) else [v]
                n += sum(1 for x in vals if x is not None and x != "" or isinstance(x, ExcelError))
            return float(n)
        if name == "IF":
            if not 1 <= len(args) <= 3:
                raise UnsupportedFormula("IF takes 1-3 arguments")
            cond = to_bool(self._scalar(self.ev(args[0], sheet)))
            if isinstance(cond, ExcelError):
                return cond
            if cond:
                return self._branch(args[1] if len(args) > 1 else ("lit", True), sheet)
            return self._branch(args[2], sheet) if len(args) > 2 else False
        if name == "IFERROR":
            self._arity(name, args, 2, 2)
            v = self._scalar(self.ev(args[0], sheet))
            return self._branch(args[1], sheet) if isinstance(v, ExcelError) else v
        if name == "VLOOKUP":
            return self._vlookup(args, sheet)
        if name == "ROUND":
            self._arity(name, args, 2, 2)
            x, n = (to_number(self._scalar(self.ev(a, sheet))) for a in args)
            for v in (x, n):
                if isinstance(v, ExcelError):
                    return v
            return round_half_away(x, int(n))
        if name == "ABS":
            self._arity(name, args, 1, 1)
            x = to_number(self._scalar(self.ev(args[0], sheet)))
            return x if isinstance(x, ExcelError) else abs(x)
        if name in ("AND", "OR"):
            vals: list[bool] = []
            for a in args:
                v = self.ev(a, sheet)
                items = [x for x in v.flat() if isinstance(x, (bool, int, float, ExcelError))] if isinstance(v, Range) else [v]
                for x in items:
                    b = to_bool(x)
                    if isinstance(b, ExcelError):
                        return b
                    vals.append(b)
            if not vals:
                return VALUE
            return all(vals) if name == "AND" else any(vals)
        if name == "NOT":
            self._arity(name, args, 1, 1)
            b = to_bool(self._scalar(self.ev(args[0], sheet)))
            return b if isinstance(b, ExcelError) else not b
        if name in ("CONCATENATE", "CONCAT"):
            parts = []
            for a in args:
                v = self.ev(a, sheet)
                for x in (v.flat() if isinstance(v, Range) else [v]):
                    t = to_text(x)
                    if isinstance(t, ExcelError):
                        return t
                    parts.append(t)
            return "".join(parts)
        raise UnsupportedFormula(f"function {name} is not supported")

    def _branch(self, node: Any, sheet: str) -> Any:
        v = self.ev(node, sheet)
        if isinstance(v, Range):
            v = self._scalar(v)
        return 0.0 if v is None and node[0] == "lit" else v

    @staticmethod
    def _arity(name: str, args: list[Any], lo: int, hi: int) -> None:
        if not lo <= len(args) <= hi:
            raise UnsupportedFormula(f"{name} takes {lo}-{hi} arguments")

    def _vlookup(self, args: list[Any], sheet: str) -> Any:
        self._arity("VLOOKUP", args, 3, 4)
        key = self._scalar(self.ev(args[0], sheet))
        table = self.ev(args[1], sheet)
        col = to_number(self._scalar(self.ev(args[2], sheet)))
        approx: Any = True
        if len(args) == 4:
            approx = to_bool(self._scalar(self.ev(args[3], sheet)))
        for v in (key, col, approx):
            if isinstance(v, ExcelError):
                return v
        if not isinstance(table, Range):
            return NA
        col_i = int(col)
        if col_i < 1:
            return VALUE
        if not table.rows or col_i > len(table.rows[0]):
            return REF
        if key is None:
            key = 0.0
        if not approx:
            for row in table.rows:
                first = row[0]
                if first is not None and _type_rank(first) == _type_rank(key) and compare(first, key) == 0:
                    return row[col_i - 1]
                if isinstance(key, str) and isinstance(first, str) and _wild_eq(key, first):
                    return row[col_i - 1]
            return NA
        # approximate: last row whose first value <= key (data assumed ascending, same type)
        found = None
        for row in table.rows:
            first = row[0]
            if first is None or _type_rank(first) != _type_rank(key):
                continue
            if compare(first, key) <= 0:
                found = row
            else:
                break
        return NA if found is None else found[col_i - 1]


def _wild_eq(pattern: str, value: str) -> bool:
    """Exact-match VLOOKUP supports * ? and ~ wildcards for text keys."""
    if not any(ch in pattern for ch in "*?"):
        return False
    rx, i = [], 0
    while i < len(pattern):
        ch = pattern[i]
        if ch == "~" and i + 1 < len(pattern):
            rx.append(re.escape(pattern[i + 1]))
            i += 2
            continue
        rx.append(".*" if ch == "*" else "." if ch == "?" else re.escape(ch))
        i += 1
    return re.fullmatch("".join(rx), value, re.IGNORECASE | re.DOTALL) is not None


# ---- workbook binding ------------------------------------------------------------------------


def cell_formula(cell: Any) -> str | None:
    """The formula text of an openpyxl cell (normal, array or data-table), else None."""
    v = cell.value
    if isinstance(v, str) and cell.data_type == "f":
        return v
    text = getattr(v, "text", None)  # ArrayFormula
    if isinstance(text, str):
        return text if text.startswith("=") else "=" + text
    if type(v).__name__ == "DataTableFormula":
        return "=TABLE()"
    return None


def native(v: Any) -> Any:
    """openpyxl cell value -> evaluator value (dates become Excel serial numbers)."""
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time, _dt.timedelta)):
        from openpyxl.utils.datetime import to_excel

        return float(to_excel(v))
    if isinstance(v, str) and v in ERROR_CODES:
        return ExcelError(v)
    if isinstance(v, int) and not isinstance(v, bool):
        return float(v)
    return v


class WorkbookEvaluator:
    """Evaluate cells of an openpyxl workbook loaded with formulas.

    ``values_wb`` (the same file loaded with data_only=True) supplies cached results; when
    ``prefer_cached`` is true, referenced formula cells use their cached value if present.
    """

    def __init__(self, formulas_wb: Any, values_wb: Any | None = None, prefer_cached: bool = True):
        self.fwb, self.vwb, self.prefer_cached = formulas_wb, values_wb, prefer_cached
        self._memo: dict[tuple[str, int, int], Any] = {}
        self._active: set[tuple[str, int, int]] = set()
        self.ev = Evaluator(self._get, formulas_wb.sheetnames, lambda s: formulas_wb[s].max_row)

    def cached(self, sheet: str, row: int, col: int) -> Any:
        if self.vwb is None:
            return None
        return self.vwb[sheet].cell(row=row, column=col).value

    def _get(self, sheet: str, row: int, col: int) -> Any:
        key = (sheet, row, col)
        if key in self._memo:
            return self._memo[key]
        ws = self.fwb[sheet]
        if row > ws.max_row or col > ws.max_column:
            return None
        cell = ws.cell(row=row, column=col)
        f = cell_formula(cell)
        if f is None:
            v = native(cell.value)
        else:
            cached = self.cached(sheet, row, col) if self.prefer_cached else None
            v = native(cached) if cached is not None else self.evaluate(sheet, row, col)
        self._memo[key] = v
        return v

    def evaluate(self, sheet: str, row: int, col: int) -> Any:
        """Evaluate the formula in a cell (ignores its own cached value)."""
        key = (sheet, row, col)
        cell = self.fwb[sheet].cell(row=row, column=col)
        f = cell_formula(cell)
        if f is None:
            return native(cell.value)
        if key in self._active:
            raise UnsupportedFormula(f"circular reference at {sheet}!{cell.coordinate}")
        self._active.add(key)
        try:
            return self.ev.eval_formula(f, sheet)
        finally:
            self._active.discard(key)
